#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown to Word converter for A35 Linux report content.
Generates a .docx with headings, tables, blockquotes and paragraphs.
"""

import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


def set_chinese_font(run, font_name='SimSun', size=10.5, bold=False):
    """Set Chinese font for a run."""
    font = run.font
    font.name = font_name
    font.size = Pt(size)
    font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def add_heading(doc, text, level):
    """Add heading with Chinese font."""
    sizes = {0: 18, 1: 16, 2: 14, 3: 12}
    size = sizes.get(level, 12)
    p = doc.add_heading(level=min(level + 1, 9))
    run = p.add_run(text)
    set_chinese_font(run, 'SimHei', size, bold=True)
    return p


def add_paragraph(doc, text, bold=False, italic=False, size=10.5, alignment=None):
    """Add normal paragraph with Chinese font."""
    p = doc.add_paragraph()
    if alignment:
        p.alignment = alignment
    run = p.add_run(text)
    set_chinese_font(run, 'SimSun', size, bold=bold)
    run.italic = italic
    # Set line spacing
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    return p


def parse_table(lines, start_idx):
    """Parse markdown table, return (rows, end_idx)."""
    rows = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        row_text = lines[i].strip()
        # Skip separator line like |---|---|
        if re.match(r'^\|(\s*[-:]+\s*\|)+$', row_text):
            i += 1
            continue
        cells = [cell.strip() for cell in row_text.split('|')[1:-1]]
        rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    """Add table to document."""
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'
    for i, row in enumerate(rows):
        cells = table.rows[i].cells
        for j, cell_text in enumerate(row):
            if j >= len(cells):
                break
            cells[j].text = cell_text
            # Set font for all paragraphs in cell
            for paragraph in cells[j].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_chinese_font(run, 'SimSun', 10)
                if i == 0:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph()


def md_to_docx(md_path, docx_path):
    doc = Document()
    # Set default document font
    style = doc.styles['Normal']
    style.font.name = 'SimSun'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
    style.font.size = Pt(10.5)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_quote = False
    quote_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            if in_quote and quote_lines:
                # End quote block
                add_paragraph(doc, '\n'.join(quote_lines), italic=True, size=10)
                quote_lines = []
                in_quote = False
            i += 1
            continue

        # Blockquote (AI prompt)
        if stripped.startswith('> '):
            in_quote = True
            quote_lines.append(stripped[2:])
            i += 1
            continue
        elif in_quote:
            add_paragraph(doc, '\n'.join(quote_lines), italic=True, size=10)
            quote_lines = []
            in_quote = False

        # Headings
        if stripped.startswith('# '):
            add_heading(doc, stripped[2:], 0)
            i += 1
            continue
        elif stripped.startswith('## '):
            add_heading(doc, stripped[3:], 1)
            i += 1
            continue
        elif stripped.startswith('### '):
            add_heading(doc, stripped[4:], 2)
            i += 1
            continue
        elif stripped.startswith('#### '):
            add_heading(doc, stripped[5:], 3)
            i += 1
            continue

        # Tables
        if stripped.startswith('|') and not stripped.startswith('| 字段 |'):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        # Horizontal rules
        if stripped == '---':
            i += 1
            continue

        # Bold text markers
        text = stripped
        # Replace markdown bold
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        # Replace inline code
        text = re.sub(r'`(.*?)`', r'\1', text)

        # Image/table placeholders
        if text.startswith('[图 ') or text.startswith('[表 '):
            add_paragraph(doc, text, bold=True, size=11, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        elif text.startswith('>') or text.startswith('**AI 绘图提示词'):
            pass
        else:
            add_paragraph(doc, text)

        i += 1

    # Flush remaining quote
    if quote_lines:
        add_paragraph(doc, '\n'.join(quote_lines), italic=True, size=10)

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == '__main__':
    md_path = '/home/alientek/dvr_project/mier/文档/A35_Linux_侧报告内容草稿.md'
    docx_path = '/home/alientek/dvr_project/mier/文档/A35_Linux_侧报告内容.docx'
    md_to_docx(md_path, docx_path)
